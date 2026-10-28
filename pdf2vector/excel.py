"""Geração da planilha Excel (.xlsx) com tudo o que foi extraído e vetorizado."""

from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .extractor import DocumentData
from .vectorizer import Vetorizacao

LIMITE_CELULA = 32767  # máximo de caracteres por célula no Excel
MAX_LINHAS = 1_048_575

_CAB_FONTE = Font(bold=True, color="FFFFFF")
_CAB_FUNDO = PatternFill("solid", fgColor="1F4E78")
_TITULO = Font(bold=True, size=14)
_SUBTITULO = Font(bold=True, size=12)
_NEGRITO = Font(bold=True)
_QUEBRA = Alignment(wrap_text=True, vertical="top")

COLUNAS_BBOX = ["x0", "y0", "x1", "y1"]


def _limpo(v):
    if isinstance(v, str):
        v = ILLEGAL_CHARACTERS_RE.sub("", v)
        if len(v) > LIMITE_CELULA:
            v = v[: LIMITE_CELULA - 15] + " …[truncado]"
    elif isinstance(v, (list, tuple, dict)):
        v = _limpo(str(v))
    return v


_ACENTOS = {"pagina": "página", "indice": "índice", "rotacao": "rotação", "direcao": "direção",
            "nivel": "nível", "titulo": "título", "conteudo": "conteúdo", "ocorrencias": "ocorrências",
            "rotulo": "rótulo", "ancora": "âncora", "anotacoes": "anotações", "sumario": "sumário",
            "e": "é", "dimensao": "dimensão", "caracteres": "caracteres", "criptografado": "criptografado",
            "formulario": "formulário", "chave": "chave", "tamanho": "tamanho"}


def _rotulo(chave: str) -> str:
    return " ".join(_ACENTOS.get(p, p) for p in chave.split("_")).capitalize()


def _cab(ws, valores) -> None:
    linha = []
    for v in valores:
        cell = WriteOnlyCell(ws, _limpo(v))
        cell.font, cell.fill = _CAB_FONTE, _CAB_FUNDO
        linha.append(cell)
    ws.append(linha)


def _aba(wb: Workbook, titulo: str, registros: list[dict], colunas: list[str],
         larguras: dict[str, int] | None = None, quebra: tuple[str, ...] = ()) -> None:
    # Planilha em modo write-only: as linhas vão direto para o arquivo, sem
    # manter todas as células na memória (essencial para PDFs de milhares de páginas).
    ws = wb.create_sheet(titulo)
    larguras = larguras or {}
    for i, c in enumerate(colunas, start=1):
        ws.column_dimensions[get_column_letter(i)].width = larguras.get(c, max(10, min(18, len(c) + 4)))
    ws.freeze_panes = "A2"
    if registros:
        n = min(len(registros), MAX_LINHAS) + 1
        ws.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{n}"
    _cab(ws, [_rotulo(c) for c in colunas])
    idx_quebra = {i for i, c in enumerate(colunas) if c in quebra}
    for reg in registros[:MAX_LINHAS]:
        linha = [_limpo(reg.get(c, "")) for c in colunas]
        for i in idx_quebra:
            cell = WriteOnlyCell(ws, linha[i])
            cell.alignment = _QUEBRA
            linha[i] = cell
        ws.append(linha)
    if len(registros) > MAX_LINHAS:
        ws.append([f"… {len(registros) - MAX_LINHAS} linhas omitidas (limite do Excel)"])


def build_workbook(doc: DocumentData, vet: Vetorizacao | None = None,
                   *, incluir_trechos: bool = True) -> Workbook:
    wb = Workbook(write_only=True)
    _resumo(wb.create_sheet("Resumo"), doc, vet)

    _aba(wb, "Metadados", [{"campo": _rotulo(k), "valor": v} for k, v in doc.metadados.items()],
         ["campo", "valor"], {"campo": 28, "valor": 70})

    _aba(wb, "Páginas", doc.paginas,
         ["pagina", "largura_pt", "altura_pt", "largura_mm", "altura_mm", "rotacao", "caracteres",
          "palavras", "blocos_texto", "tabelas", "imagens", "links", "desenhos_vetoriais", "ocr",
          "sem_texto", "texto"],
         {"texto": 100}, quebra=("texto",))

    _aba(wb, "Blocos de texto", doc.blocos,
         ["pagina", "bloco", "linhas", "texto", *COLUNAS_BBOX], {"texto": 90}, quebra=("texto",))

    _aba(wb, "Linhas", doc.linhas,
         ["pagina", "bloco", "linha", "texto", *COLUNAS_BBOX, "direcao"], {"texto": 80})

    if incluir_trechos:
        _aba(wb, "Formatação", doc.trechos,
             ["pagina", "bloco", "linha", "trecho", "texto", "fonte", "tamanho", "estilo", "cor",
              *COLUNAS_BBOX], {"texto": 60, "fonte": 28})

    _aba(wb, "Fontes", doc.fontes, ["fonte", "tamanho", "estilo", "ocorrencias", "caracteres"],
         {"fonte": 36})

    _tabelas(wb, doc)

    _aba(wb, "Imagens", doc.imagens,
         ["pagina", "indice", "xref", "largura_px", "altura_px", "bits_por_componente", "espaco_cor",
          "dpi_x", "dpi_y", "bytes", "md5", *COLUNAS_BBOX], {"md5": 34})

    _aba(wb, "Desenhos vetoriais", doc.desenhos,
         ["pagina", "indice", "tipo", "elementos", "cor_contorno", "cor_preenchimento", "espessura",
          "fechado", *COLUNAS_BBOX, "pontos"], {"elementos": 30, "pontos": 80})

    _aba(wb, "Links", doc.links, ["pagina", "tipo", "destino", "texto_ancora", *COLUNAS_BBOX],
         {"destino": 50, "texto_ancora": 40})
    _aba(wb, "Anotações", doc.anotacoes,
         ["pagina", "tipo", "conteudo", "autor", "assunto", "criada_em", *COLUNAS_BBOX],
         {"conteudo": 60})
    _aba(wb, "Formulário", doc.campos, ["pagina", "nome", "tipo", "valor", "rotulo", *COLUNAS_BBOX],
         {"nome": 30, "valor": 40})
    _aba(wb, "Sumário", doc.sumario, ["nivel", "titulo", "pagina"], {"titulo": 70})

    if vet is not None:
        _vetores(wb, vet)

    if doc.avisos:
        _aba(wb, "Avisos", [{"aviso": a} for a in doc.avisos], ["aviso"], {"aviso": 120})
    return wb


def _resumo(ws, doc: DocumentData, vet: Vetorizacao | None) -> None:
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 60
    titulo = WriteOnlyCell(ws, _limpo(f"Relatório do PDF: {doc.arquivo}"))
    titulo.font = _TITULO
    ws.append([titulo])
    ws.append([f"Gerado em {datetime.now():%d/%m/%Y %H:%M} por pdf2vector"])
    ws.append([])
    m = doc.metadados
    linhas = [
        ("Páginas", m.get("paginas")),
        ("Título", m.get("titulo")),
        ("Autor", m.get("autor")),
        ("Palavras", m.get("total_palavras")),
        ("Caracteres", m.get("total_caracteres")),
        ("Tabelas detectadas", m.get("total_tabelas")),
        ("Imagens", m.get("total_imagens")),
        ("Links", m.get("total_links")),
        ("Desenhos vetoriais", m.get("total_desenhos_vetoriais")),
        ("Anotações", len(doc.anotacoes)),
        ("Campos de formulário", len(doc.campos)),
        ("Entradas no sumário", len(doc.sumario)),
        ("Páginas sem texto (possível digitalização)",
         sum(1 for p in doc.paginas if p["sem_texto"] == "sim")),
    ]
    if vet is not None:
        linhas += [
            ("Trechos vetorizados (chunks)", len(vet.chunks)),
            ("Motor de vetorização", vet.motor),
            ("Modelo", vet.modelo),
            ("Dimensão dos vetores", vet.dimensao),
        ]
    if doc.avisos:
        linhas.append(("Avisos", f"{len(doc.avisos)} (veja a aba Avisos)"))
    for k, v in linhas:
        rotulo = WriteOnlyCell(ws, k)
        rotulo.font = _NEGRITO
        ws.append([rotulo, _limpo(v if v is not None else "")])


def _tabelas(wb: Workbook, doc: DocumentData) -> None:
    """Uma aba-índice com todas as tabelas, empilhadas e identificadas, e uma
    aba própria para cada tabela (até 200), pronta para filtrar e calcular."""
    ws = wb.create_sheet("Tabelas")
    if not doc.tabelas:
        ws.append(["Nenhuma tabela detectada no documento."])
        return
    for t in doc.tabelas:
        titulo = WriteOnlyCell(ws, f"Tabela {t['id']} — página {t['pagina']} "
                                   f"({t['linhas']} linhas × {t['colunas']} colunas)")
        titulo.font = _SUBTITULO
        ws.append([titulo])
        if t["cabecalho"]:
            _cab(ws, t["cabecalho"])
            dados = t["dados"][1:] if t["dados"] and t["dados"][0] == t["cabecalho"] else t["dados"]
        else:
            dados = t["dados"]
        for row in dados:
            ws.append([_celula(ws, v) for v in row])
        ws.append([])

    for t in doc.tabelas[:200]:
        aba = wb.create_sheet(f"Tab{t['id']} p{t['pagina']}"[:31])
        for j in range(1, (t["colunas"] or 0) + 1):
            aba.column_dimensions[get_column_letter(j)].width = 18
        aba.freeze_panes = "A2"
        for i, row in enumerate(t["dados"]):
            if i == 0:
                _cab(aba, row)
            else:
                aba.append([_celula(aba, v) for v in row])


def _celula(ws, v):
    valor = _converter(v)
    if isinstance(valor, float) and isinstance(v, str) and v.strip().endswith("%"):
        cell = WriteOnlyCell(ws, valor)
        cell.number_format = "0.0%"
        return cell
    return valor


def _converter(v):
    """Transforma números escritos no padrão brasileiro (1.234,56 / 12%) em
    número de verdade, para que somas e filtros funcionem no Excel."""
    if not isinstance(v, str):
        return v
    s = v.strip().replace("R$", "").replace(" ", "").strip()
    if not s or len(s) > 30:
        return _limpo(v)
    pct = s.endswith("%")
    if pct:
        s = s[:-1].strip()
    neg = s.startswith("(") and s.endswith(")")
    if neg:
        s = s[1:-1]
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    elif s.count(".") > 1:
        s = s.replace(".", "")
    if len(s.lstrip("-")) > 1 and s.lstrip("-")[0] == "0" and s.lstrip("-")[1] != ".":
        return _limpo(v)  # códigos com zero à esquerda (CEP, matrícula) ficam como texto
    try:
        num = float(s)
    except ValueError:
        return _limpo(v)
    if not s.lstrip("-").replace(".", "").isdigit():
        return _limpo(v)
    if neg:
        num = -num
    if pct:
        num /= 100
    return int(num) if num.is_integer() and not pct and "." not in s else num


def _vetores(wb: Workbook, vet: Vetorizacao) -> None:
    _aba(wb, "Chunks", [c.__dict__ for c in vet.chunks],
         ["id", "pagina_inicial", "pagina_final", "palavras", "caracteres", "palavras_chave", "texto"],
         {"palavras_chave": 40, "texto": 100}, quebra=("texto",))

    ws = wb.create_sheet("Vetores")
    ws.freeze_panes = "D2"
    _cab(ws, ["Chunk", "Página inicial", "Página final"] + [f"d{i}" for i in range(vet.dimensao)])
    for c, v in zip(vet.chunks, vet.vetores):
        ws.append([c.id, c.pagina_inicial, c.pagina_final] + [round(float(x), 6) for x in v])


def write_excel(doc: DocumentData, vet: Vetorizacao | None, destino: str | Path | None = None,
                **kwargs) -> bytes:
    """Gera o .xlsx. Se ``destino`` for informado, grava o arquivo; sempre
    devolve os bytes (útil para respostas HTTP)."""
    wb = build_workbook(doc, vet, **kwargs)
    buf = io.BytesIO()
    wb.save(buf)
    dados = buf.getvalue()
    if destino is not None:
        Path(destino).write_bytes(dados)
    return dados
