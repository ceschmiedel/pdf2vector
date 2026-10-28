import importlib.util
from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook

from pdf2vector import converter, extract, vectorize
from pdf2vector.cli import main
from pdf2vector.excel import _converter
from pdf2vector.vectorizer import chunk_text

RAIZ = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="session")
def pdf(tmp_path_factory):
    spec = importlib.util.spec_from_file_location("gerar", RAIZ / "exemplos" / "gerar_pdf_exemplo.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    destino = tmp_path_factory.mktemp("pdf") / "exemplo.pdf"
    mod.gerar(str(destino))
    return destino


def test_extrai_conteudo_completo(pdf):
    doc = extract(pdf)
    assert doc.metadados["paginas"] == 2
    assert doc.metadados["titulo"] == "Relatório Trimestral"
    assert "café" in doc.paginas[0]["texto"]
    assert any(t["estilo"] == "negrito" for t in doc.trechos)
    assert len(doc.tabelas) == 1
    tab = doc.tabelas[0]
    assert tab["pagina"] == 1 and tab["colunas"] == 3
    assert ["Sul", "1.234,56", "12%"] in tab["dados"]
    assert len(doc.imagens) == 1 and doc.imagens[0]["largura_px"] == 40
    assert doc.links[0]["destino"] == "https://example.com/vendas"
    assert doc.anotacoes[0]["conteudo"] == "Conferir números do Nordeste"
    assert doc.campos[0]["nome"] == "responsavel" and doc.campos[0]["valor"] == "Maria Souza"
    assert [s["titulo"] for s in doc.sumario] == ["Relatório Trimestral de Vendas", "Conclusões"]
    assert any("curva" in d["elementos"] for d in doc.desenhos)  # o círculo


def test_chunks_cobrem_texto_com_sobreposicao():
    paginas = [" ".join(f"p1w{i}" for i in range(150)), " ".join(f"p2w{i}" for i in range(150))]
    chunks = chunk_text(paginas, tamanho=100, sobreposicao=20)
    assert [c.palavras for c in chunks] == [100, 100, 100, 60]
    assert chunks[0].pagina_inicial == 1 and chunks[1].pagina_final == 2
    assert chunks[0].texto.split()[-20:] == chunks[1].texto.split()[:20]
    assert chunks[-1].texto.split()[-1] == "p2w149"


def test_vetores_normalizados_e_busca(pdf):
    doc = extract(pdf)
    vet = vectorize(doc.texto_por_pagina, tamanho_chunk=40, sobreposicao=5)
    assert vet.vetores.shape == (len(vet.chunks), 256)
    assert np.allclose(np.linalg.norm(vet.vetores, axis=1), 1, atol=1e-5)
    melhor, score = vet.buscar("distribuição no Nordeste custos logísticos", k=1)[0]
    assert melhor.pagina_inicial == 2 and score > 0.3


def test_documento_vazio_nao_quebra():
    vet = vectorize(["", "   "])
    assert vet.chunks == [] and vet.vetores.shape == (0, 256)


def test_excel_tem_todas_as_abas(pdf, tmp_path):
    destino = tmp_path / "saida.xlsx"
    converter(pdf, destino, tamanho_chunk=40)
    wb = load_workbook(destino)
    for aba in ["Resumo", "Metadados", "Páginas", "Blocos de texto", "Linhas", "Formatação",
                "Fontes", "Tabelas", "Tab1 p1", "Imagens", "Desenhos vetoriais", "Links",
                "Anotações", "Formulário", "Sumário", "Chunks", "Vetores"]:
        assert aba in wb.sheetnames, aba
    tab = wb["Tab1 p1"]
    assert tab["A1"].value == "Região"
    assert tab["B2"].value == pytest.approx(1234.56)
    assert tab["C3"].value == pytest.approx(-0.035)
    vet = wb["Vetores"]
    assert vet.max_column == 3 + 256
    assert vet.max_row == wb["Chunks"].max_row


@pytest.mark.parametrize("entrada,esperado", [
    ("1.234,56", 1234.56), ("R$ 10,00", 10.0), ("12%", 0.12), ("(5,5)", -5.5),
    ("2026", 2026), ("01310-100", "01310-100"), ("0042", "0042"), ("abc", "abc"),
    ("1e5", "1e5"), ("123.456.789-00", "123.456.789-00"),
])
def test_conversao_numeros_brasileiros(entrada, esperado):
    assert _converter(entrada) == esperado


def test_cli(pdf, tmp_path, capsys):
    destino = tmp_path / "cli.xlsx"
    assert main([str(pdf), "-o", str(destino), "--chunk", "50"]) == 0
    assert destino.exists()
    assert "2 páginas" in capsys.readouterr().out


def test_web(pdf):
    pytest.importorskip("fastapi")
    pytest.importorskip("httpx")
    from fastapi.testclient import TestClient
    from pdf2vector.web import app

    c = TestClient(app)
    assert c.get("/").status_code == 200
    r = c.post("/converter", files={"arquivo": ("relatório.pdf", pdf.read_bytes(), "application/pdf")})
    assert r.status_code == 200
    assert r.headers["x-paginas"] == "2"
    assert r.content[:2] == b"PK"
    r = c.post("/converter", files={"arquivo": ("x.pdf", b"nada", "application/pdf")})
    assert r.status_code == 415
